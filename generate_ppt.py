# -*- coding: utf-8 -*-
"""
答辩PPT生成脚本
桂林电子科技大学 网络工程专业
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Cm
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import os

def create_ppt():
    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 16:9
    prs.slide_height = Inches(7.5)

    # 颜色主题
    PRIMARY = RGBColor(0x4A, 0x90, 0xD9)
    DARK = RGBColor(0x2C, 0x3E, 0x50)
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    LIGHT_BG = RGBColor(0xF5, 0xF7, 0xFA)
    GRAY = RGBColor(0x99, 0x99, 0x99)
    GREEN = RGBColor(0x52, 0xC4, 0x1A)

    def add_bg(slide, color=LIGHT_BG):
        """设置幻灯片背景"""
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_title_bar(slide, title_text, subtitle_text=''):
        """添加顶部标题栏"""
        # 蓝色顶部条
        from pptx.util import Inches
        shape = slide.shapes.add_shape(
            1, Inches(0), Inches(0), prs.slide_width, Inches(1.2)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = PRIMARY
        shape.line.fill.background()

        tf = shape.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run()
        r.text = title_text
        r.font.size = Pt(36)
        r.font.color.rgb = WHITE
        r.font.bold = True
        tf.margin_left = Inches(0.8)

    def add_text_box(slide, left, top, width, height, text, font_size=24, color=DARK, bold=False, alignment=PP_ALIGN.LEFT):
        """添加文本框"""
        txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = alignment
        r = p.add_run()
        r.text = text
        r.font.size = Pt(font_size)
        r.font.color.rgb = color
        r.font.bold = bold
        return tf

    def add_bullet_points(slide, left, top, width, height, points, font_size=22):
        """添加要点列表"""
        txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        tf = txBox.text_frame
        tf.word_wrap = True
        for i, point in enumerate(points):
            if i == 0:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()
            p.space_after = Pt(12)
            r = p.add_run()
            r.text = f'● {point}'
            r.font.size = Pt(font_size)
            r.font.color.rgb = DARK
        return tf

    # ═══════════════════════════
    # Slide 1: 封面
    # ═══════════════════════════
    slide1 = prs.slides.add_slide(prs.slide_layouts[6])  # 空白布局
    # 蓝色背景
    bg = slide1.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = PRIMARY

    add_text_box(slide1, 1.5, 1.5, 10, 1.5,
                 '基于微信小程序的英语学习打卡系统\n设计与实现',
                 font_size=44, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)

    add_text_box(slide1, 1.5, 3.2, 10, 0.8,
                 'Design and Implementation of English Learning\nCheck-in System Based on WeChat Mini Program',
                 font_size=20, color=RGBColor(0xE0, 0xE8, 0xF0), alignment=PP_ALIGN.CENTER)

    add_text_box(slide1, 1.5, 4.8, 10, 2,
                 '答辩人：李新国\n专  业：网络工程（2023级）\n指导老师：朱凯\n桂林电子科技大学 · 计算机与信息安全学院',
                 font_size=24, color=WHITE, alignment=PP_ALIGN.CENTER, bold=False)

    # ═══════════════════════════
    # Slide 2: 目录
    # ═══════════════════════════
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide2)
    add_title_bar(slide2, '汇报提纲')

    contents = [
        ('01', '项目背景与意义'),
        ('02', '系统总体设计'),
        ('03', '核心技术实现'),
        ('04', '网络安全特色模块'),
        ('05', '系统演示'),
        ('06', '测试与总结'),
    ]
    for i, (num, title) in enumerate(contents):
        y = 1.8 + i * 0.85
        # 数字
        add_text_box(slide2, 1.5, y, 0.8, 0.6, num, font_size=32, color=PRIMARY, bold=True)
        # 标题
        add_text_box(slide2, 2.3, y, 8, 0.6, title, font_size=28, color=DARK)

    # ═══════════════════════════
    # Slide 3: 项目背景
    # ═══════════════════════════
    slide3 = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide3)
    add_title_bar(slide3, '01  项目背景与意义')

    add_bullet_points(slide3, 0.8, 1.6, 5.5, 5, [
        '移动互联网时代，碎片化学习成为大学生自主学习主要方式',
        '英语学习具有长期性、重复性、积累性特点',
        '传统纸质记录和通用APP存在学习难监督、数据难记录等痛点',
        '微信小程序"无需下载、即用即走"，用户基础庞大（12亿+月活）',
        '学生群体微信渗透率接近100%，天然适合校园应用场景',
    ], font_size=22)

    add_bullet_points(slide3, 7, 1.6, 5.5, 5, [
        '本系统的创新点：',
        '后端采用Python Django + DRF轻量框架',
        '深度融合网络工程专业知识——构建四层安全防护体系',
        'IP溯源 + 异常登录检测 + JWT鉴权 + HTTPS加密',
        '紧贴高校学生实际需求：多维度排行榜、打卡日历、进度可视化',
        '同期同类系统多以Java为主，本系统Python方案开发效率更优',
    ], font_size=22)

    # ═══════════════════════════
    # Slide 4: 系统架构
    # ═══════════════════════════
    slide4 = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide4)
    add_title_bar(slide4, '02  系统总体架构设计')

    add_text_box(slide4, 0.8, 1.6, 11, 0.6,
                 '系统采用前后端分离三层架构：表现层（小程序）→ 业务逻辑层（Django）→ 数据访问层（MySQL）',
                 font_size=22, color=DARK, bold=True)

    # 架构三层
    layers = [
        ('表现层', '微信小程序\n(WXML + WXSS + JS)', '用户交互 · 页面渲染 · wx.request API调用'),
        ('业务逻辑层', 'Python Django 4.2\n+ Django REST Framework', 'RESTful API · JWT认证 · 业务逻辑 · 安全中间件'),
        ('数据访问层', 'MySQL 8.0\n+ Django ORM', '数据持久化 · 外键约束 · 索引优化 · 事务管理'),
    ]
    for i, (name, tech, desc) in enumerate(layers):
        y = 2.5 + i * 1.4
        # 层名称卡片
        shape = slide4.shapes.add_shape(
            1, Inches(0.8), Inches(y), Inches(2), Inches(1.1)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = PRIMARY
        shape.line.fill.background()
        tf = shape.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = name
        r.font.size = Pt(22)
        r.font.color.rgb = WHITE
        r.font.bold = True

        add_text_box(slide4, 3.2, y, 4, 1.1, tech, font_size=20, color=DARK)
        add_text_box(slide4, 7.5, y, 5.5, 1.1, desc, font_size=18, color=GRAY)

    # ═══════════════════════════
    # Slide 5: 功能模块
    # ═══════════════════════════
    slide5 = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide5)
    add_title_bar(slide5, '02  功能模块与数据库设计')

    modules = [
        ('用户模块', '微信登录、个人中心\n学习统计'),
        ('单词模块', '单词浏览/搜索\n学习进度追踪'),
        ('打卡模块', '每日打卡、连续天数\n打卡日历'),
        ('排行榜模块', '日/周/月/总榜\n多维度排名'),
        ('管理后台', '用户/单词管理\n数据统计可视化'),
        ('安全模块', 'IP溯源、异常检测\nJWT鉴权、HTTPS'),
    ]
    for i, (name, desc) in enumerate(modules):
        col = i % 3
        row = i // 3
        x = 0.6 + col * 4.2
        y = 1.5 + row * 2.8
        shape = slide5.shapes.add_shape(
            1, Inches(x), Inches(y), Inches(3.8), Inches(2.4)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shape.line.color.rgb = PRIMARY
        shape.line.width = Pt(2)

        tf = shape.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = name
        r.font.size = Pt(28)
        r.font.color.rgb = PRIMARY
        r.font.bold = True

        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        p2.space_before = Pt(12)
        r2 = p2.add_run()
        r2.text = desc
        r2.font.size = Pt(18)
        r2.font.color.rgb = DARK

    # ═══════════════════════════
    # Slide 6: 核心技术
    # ═══════════════════════════
    slide6 = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide6)
    add_title_bar(slide6, '03  核心技术实现 — Django后端')

    code_blocks = [
        ('微信登录流程', '小程序wx.login() → 后端获取code → 微信接口换取openid\n→ 自动注册/登录 → 生成JWT Token → 记录登录IP日志'),
        ('打卡核心逻辑', '检查今日是否已打卡 → 查询昨日记录 → 计算连续天数\n→ 创建打卡记录 → 更新用户统计 → 事务保证一致性'),
        ('排行榜算法', '按时间维度筛选打卡记录 → 按用户分组聚合单词数\n→ 降序排列取Top50 → 同时返回个人排名信息'),
        ('管理后台', 'Dashboard聚合查询（总用户/今日打卡/累计单词/异常登录）\n→ 用户列表管理 → 登录日志审查 → 用户状态切换'),
    ]
    for i, (title, desc) in enumerate(code_blocks):
        y = 1.5 + i * 1.4
        add_text_box(slide6, 0.8, y, 3, 0.5, title, font_size=22, color=PRIMARY, bold=True)
        add_text_box(slide6, 0.8, y + 0.4, 11, 1, desc, font_size=18, color=DARK)

    # ═══════════════════════════
    # Slide 7: 安全模块
    # ═══════════════════════════
    slide7 = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide7)
    add_title_bar(slide7, '04  网络安全特色模块（核心亮点）')

    add_text_box(slide7, 0.8, 1.5, 11, 0.6,
                 '深度融合网络工程专业知识，构建四层安全防护体系',
                 font_size=24, color=PRIMARY, bold=True)

    security = [
        ('第一层\nIP溯源', '自定义Django中间件LoginIPTraceMiddleware\n自动提取X-Forwarded-For/X-Real-IP\n穿透代理获取客户端真实IP\n每次登录自动记录IP+设备+时间'),
        ('第二层\n异常检测', '对比当前IP与历史常用IP集合\n新增IP+历史≥2常用IP → 标记异常\n管理后台高亮显示异常登录记录\n参考主流平台"异地提醒"机制'),
        ('第三层\nJWT鉴权', 'SimpleJWT无状态认证方案\nToken包含user_id+openid\n2小时过期+7天Refresh轮转\n"默认拒绝"策略防接口越权'),
        ('第四层\nHTTPS', 'Nginx反向代理 + SSL/TLS证书\n仅启用TLSv1.2/1.3安全协议\nHTTP→HTTPS 301强制跳转\n防校园网中间人攻击'),
    ]
    for i, (name, desc) in enumerate(security):
        x = 0.5 + i * 3.15
        # 标题
        shape = slide7.shapes.add_shape(
            1, Inches(x), Inches(2.3), Inches(2.9), Inches(1.1)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = PRIMARY
        shape.line.fill.background()
        tf = shape.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = name
        r.font.size = Pt(20)
        r.font.color.rgb = WHITE
        r.font.bold = True

        add_text_box(slide7, x, 3.6, 2.9, 3.5, desc, font_size=16, color=DARK)

    # ═══════════════════════════
    # Slide 8: 系统演示
    # ═══════════════════════════
    slide8 = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide8)
    add_title_bar(slide8, '05  系统演示 — 核心页面')

    pages = [
        ('首页', '用户数据概览\n每日单词推荐\n快捷功能入口'),
        ('打卡页', '一键打卡按钮\n连续/累计天数\n打卡日历视图'),
        ('排行榜', '日/周/月/总榜切换\n前三名高亮显示\n我的排名展示'),
        ('个人中心', '学习数据统计\n打卡记录查看\n登录日志(IP溯源)'),
    ]
    for i, (name, desc) in enumerate(pages):
        x = 0.5 + i * 3.15
        shape = slide8.shapes.add_shape(
            1, Inches(x), Inches(1.6), Inches(2.9), Inches(5.3)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shape.line.color.rgb = PRIMARY
        shape.line.width = Pt(1.5)

        tf = shape.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = f'[{name}]'
        r.font.size = Pt(26)
        r.font.color.rgb = PRIMARY
        r.font.bold = True

        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        p2.space_before = Pt(24)
        r2 = p2.add_run()
        r2.text = desc
        r2.font.size = Pt(18)
        r2.font.color.rgb = DARK

    # ═══════════════════════════
    # Slide 9: 测试
    # ═══════════════════════════
    slide9 = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide9)
    add_title_bar(slide9, '06  系统测试与性能分析')

    add_text_box(slide9, 0.8, 1.5, 5, 0.5, '功能测试（黑盒测试）', font_size=24, color=PRIMARY, bold=True)
    add_bullet_points(slide9, 0.8, 2.1, 5.5, 2.5, [
        '覆盖全部6个核心模块',
        '微信登录、重复打卡拦截、连续打卡计算、异常登录检测等关键用例均通过',
        '边界条件测试（空数据、超长输入、并发打卡）通过',
    ], font_size=18)

    add_text_box(slide9, 0.8, 4.5, 5, 0.5, '安全性测试', font_size=24, color=PRIMARY, bold=True)
    add_bullet_points(slide9, 0.8, 5.1, 5.5, 2, [
        '未携带Token → 401拦截 ✓',
        '过期Token → 401拦截 ✓',
        'Wireshark抓包 → HTTPS加密不可读 ✓',
    ], font_size=18)

    add_text_box(slide9, 7, 1.5, 5.5, 0.5, '性能测试（JMeter压力测试）', font_size=24, color=PRIMARY, bold=True)
    add_bullet_points(slide9, 7, 2.1, 5.5, 5, [
        '50并发 → 平均响应 < 100ms',
        '100并发 → 平均响应 ~200ms',
        '200并发 → 平均响应 ~450ms（达标）',
        '500并发 → 部分接口 > 1s（可优化）',
        '结论：满足校园级应用性能需求',
    ], font_size=18)

    # ═══════════════════════════
    # Slide 10: 总结
    # ═══════════════════════════
    slide10 = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide10)
    add_title_bar(slide10, '06  总结与展望')

    add_text_box(slide10, 0.8, 1.5, 5.5, 0.5, '项目成果', font_size=26, color=PRIMARY, bold=True)
    add_bullet_points(slide10, 0.8, 2.1, 5.5, 2.5, [
        '完成微信小程序 + Django后端全栈系统',
        '实现6大功能模块，覆盖英语学习打卡全流程',
        '创新性融入4层网络安全防护体系（IP溯源/异常检测/JWT/HTTPS）',
        '系统代码总量约5000+行，API接口15+个',
    ], font_size=18)

    add_text_box(slide10, 0.8, 4.5, 5.5, 0.5, '未来展望', font_size=26, color=PRIMARY, bold=True)
    add_bullet_points(slide10, 0.8, 5.1, 5.5, 2, [
        '引入AI学习推荐算法',
        'Docker容器化 + CI/CD',
        'ELK日志分析平台集成',
    ], font_size=18)

    add_text_box(slide10, 7, 1.5, 5.5, 0.5, '技术栈总览', font_size=26, color=PRIMARY, bold=True)
    add_bullet_points(slide10, 7, 2.1, 5.5, 5, [
        '前端：微信原生小程序 + ECharts',
        '后端：Python 3.9 + Django 4.2 + DRF',
        '数据库：MySQL 8.0 + Django ORM',
        '认证：SimpleJWT 无状态认证',
        '安全：自定义中间件 + Nginx SSL',
        '测试：Postman + JMeter + Wireshark',
    ], font_size=18)

    # ═══════════════════════════
    # Slide 11: 致谢
    # ═══════════════════════════
    slide11 = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide11.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = PRIMARY

    add_text_box(slide11, 1.5, 2.5, 10, 1.5,
                 '感谢各位老师聆听！\n\n请老师批评指正',
                 font_size=48, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)

    add_text_box(slide11, 1.5, 5.5, 10, 1,
                 '桂林电子科技大学 · 计算机与信息安全学院\n网络工程专业 · 李新国',
                 font_size=20, color=RGBColor(0xE0, 0xE8, 0xF0), alignment=PP_ALIGN.CENTER)

    # ──── 保存 ────
    output_path = r'D:\毕业设计\ppt\答辩PPT_英语学习打卡系统.pptx'
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    prs.save(output_path)
    print(f'PPT已保存至: {output_path}')
    return output_path

if __name__ == '__main__':
    create_ppt()
